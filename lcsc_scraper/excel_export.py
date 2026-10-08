"""Excel 导出：按类目模式 → 每个子类目一个 sheet（参数并列为列）；
按品牌模式 → 每个品牌一个 sheet（参数合并为一列文本）。"""

from __future__ import annotations

import re
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .scraper import DATA_DIR, ScrapeTask

BASE_COLUMNS = [
    "商品编号", "型号", "品牌", "品牌网址", "品牌简介", "类目", "商品描述", "封装",
    "库存", "近期销量", "最小起订", "包装方式", "包装规格", "单价", "价格梯度",
    "毛重", "图片链接", "数据手册PDF链接", "关联(替代产品)型号", "详情链接", "简介/备注",
]

HEADER_FILL = PatternFill("solid", fgColor="1F7AE0")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    name = re.sub(r"[\\/*?:\[\]]", " ", name).strip() or "sheet"
    name = name[:28]
    candidate = name
    i = 2
    while candidate in used:
        candidate = f"{name}~{i}"
        i += 1
    used.add(candidate)
    return candidate


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _auto_width(ws, max_width: int = 46) -> None:
    for col_idx in range(1, ws.max_column + 1):
        width = 10
        for row_idx in range(1, min(ws.max_row, 60) + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                s = str(v)
                # 中文字符按2个宽度计
                w = sum(2 if ord(c) > 127 else 1 for c in s[:80])
                width = max(width, min(w + 2, max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_category_sheet(ws, products: list[dict]) -> None:
    # 参数列 = 所有商品参数名的并集，保持首次出现顺序
    param_cols: list[str] = []
    seen = set()
    for p in products:
        for k in p:
            if k.startswith("参数:") and k not in seen:
                seen.add(k)
                param_cols.append(k)
    columns = BASE_COLUMNS + param_cols

    ws.append(columns)
    _style_header(ws)
    for p in products:
        ws.append([p.get(c) for c in columns])
    _auto_width(ws)


def _param_text(p: dict) -> str:
    parts = []
    for k, v in p.items():
        if k.startswith("参数:") and v not in (None, ""):
            parts.append(f"{k[3:]}={v}")
    return " | ".join(parts)


def _write_brand_sheet(ws, products: list[dict]) -> None:
    columns = BASE_COLUMNS + ["商品参数", "参数错误"]
    ws.append(columns)
    _style_header(ws)
    for p in products:
        row = [p.get(c) for c in BASE_COLUMNS]
        row.append(_param_text(p))
        row.append(p.get("参数错误", ""))
        ws.append(row)
    _auto_width(ws)


def export_task(task: ScrapeTask) -> Path:
    DATA_DIR.mkdir(exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    site_label = {"lcsc": "立创", "hqchip": "华秋", "ickey": "云汉"}.get(task.site, task.site)
    name = (
        f"{site_label}商品数据_{task.mode}_"
        f"{'_'.join(list(task.groups)[:3])[:30]}_{stamp}.xlsx"
    )
    # 文件名去掉 Windows 非法字符
    name = re.sub(r'[\\/:*?"<>|]', "_", name)
    path = DATA_DIR / name

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for group, products in task.products.items():
        ws = wb.create_sheet(_safe_sheet_name(group, used))
        if task.mode == "category":
            _write_category_sheet(ws, products)
        else:
            _write_brand_sheet(ws, products)
    if not wb.sheetnames:
        ws = wb.create_sheet("无数据")
        ws.append(["未采集到数据"])
    wb.save(path)
    return path
